#!/usr/bin/env python3
"""
read_source.py -- turn a source file into normalized Markdown with provenance
markers so a skill can quote and cite it.

  read_source.py PATH [--out FILE] [--tables-json [FILE]] [--max-chars N] [--json]

Supported: .docx .pptx .xlsx .pdf .vtt .txt .md

Markers
  docx  every paragraph line ends with [doc p.N]; N is a running paragraph
        index (Word stores no reliable page numbers).  Heading / NW Heading
        styles become #..#### ; tables become Markdown tables [doc table N]
  pptx  ## Slide N: <title> [slide N], body text, tables, "> Notes:" lines
  xlsx  one Markdown table per sheet, capped at 200 rows, [Sheet!A1:F40]
  pdf   ## Page N [page N]
  vtt   consecutive cues by the same speaker merged: **Speaker** [HH:MM:SS]: text
  txt/md passthrough with [line N] appended every 20 lines

The header block reports the file name, type, size, extraction time, and any
Microsoft Purview sensitivity label found in docProps/custom.xml
(MSIP_Label_*_Name) or docMetadata/LabelInfo.xml (name or GUID).  It reports
only; the skill decides what to do with a labelled source.

Optional libraries: openpyxl (.xlsx), pypdf (.pdf).  A missing one is a plain
message naming the library and exit 2.
--json prints {ok, path, type, chars, label, out, tables_json}.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

TYPES = {".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx", ".pdf": "pdf",
         ".vtt": "vtt", ".txt": "txt", ".md": "md", ".markdown": "md"}
XLSX_ROW_CAP = 200
TXT_MARK_EVERY = 20

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS


class ReadError(Exception):
    def __init__(self, msg: str, hint: str = "", code: int = 1):
        super().__init__(msg)
        self.hint = hint
        self.code = code


# ---------------------------------------------------------------------------
# label detection (zip + ElementTree; no third-party code)
# ---------------------------------------------------------------------------
def read_label(path: Path) -> Optional[Dict[str, Optional[str]]]:
    """Return {"name": ..., "id": ...} when a sensitivity label is present."""
    try:
        z = zipfile.ZipFile(str(path))
    except (zipfile.BadZipFile, OSError):
        return None
    name = None
    label_id = None
    with z:
        names = set(z.namelist())
        if "docProps/custom.xml" in names:
            try:
                root = ET.fromstring(z.read("docProps/custom.xml"))
                for prop in root.iter():
                    pname = prop.get("name") or ""
                    m = re.match(r"MSIP_Label_([0-9a-fA-F-]{36})_(\w+)$", pname)
                    if not m:
                        continue
                    if m.group(2) == "Name":
                        name = "".join(t for t in prop.itertext()).strip() or None
                        label_id = label_id or m.group(1)
                    elif m.group(2) == "Enabled":
                        label_id = label_id or m.group(1)
            except ET.ParseError:
                pass
        if "docMetadata/LabelInfo.xml" in names:
            try:
                root = ET.fromstring(z.read("docMetadata/LabelInfo.xml"))
                for el in root.iter():
                    if el.tag.endswith("}label") or el.tag == "label":
                        if el.get("removed") == "1":
                            continue
                        label_id = label_id or (el.get("id") or "").strip("{}") or None
                        name = name or el.get("name")
            except ET.ParseError:
                pass
    if not name and not label_id:
        return None
    return {"name": name, "id": label_id}


def label_line(label: Optional[Dict[str, Optional[str]]]) -> Tuple[str, Optional[str]]:
    if not label:
        return "none found", None
    if label.get("name"):
        shown = label["name"]
        if label.get("id"):
            shown += " ({%s})" % label["id"]
        return shown, label["name"]
    return "{%s} (id only; the file stores no label name)" % label["id"], label["id"]


# ---------------------------------------------------------------------------
# markdown helpers
# ---------------------------------------------------------------------------
def _cell(s) -> str:
    s = "" if s is None else str(s)
    return s.replace("\r", "").replace("\n", " ").replace("|", "\\|").strip()


def md_table(header: List[str], rows: List[List[str]]) -> List[str]:
    ncol = max([len(header)] + [len(r) for r in rows]) if (header or rows) else 0
    if ncol == 0:
        return []
    header = list(header) + [""] * (ncol - len(header))
    out = ["| " + " | ".join(_cell(h) for h in header) + " |",
           "|" + "|".join([" --- "] * ncol) + "|"]
    for r in rows:
        r = list(r) + [""] * (ncol - len(r))
        out.append("| " + " | ".join(_cell(c) for c in r) + " |")
    return out


def _fmt_number(v) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    return "" if v is None else str(v)


def _need(mod: str, dist: str, for_type: str):
    try:
        return __import__(mod)
    except ImportError:
        raise ReadError("Reading %s files needs the '%s' library, which is not installed." % (for_type, dist),
                        "Run /weave:setup, or: pip install %s" % dist, code=2)


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
_HEADING_RE = re.compile(r"^(?:NW\s*)?Heading\s*(\d)$", re.IGNORECASE)


def _heading_level(style_name: str) -> Optional[int]:
    m = _HEADING_RE.match((style_name or "").strip())
    if not m:
        return None
    lvl = int(m.group(1))
    return lvl if 1 <= lvl <= 4 else 4


def _iter_block_items(container_el, doc):
    """Yield ('p', Paragraph) / ('tbl', Table) in document order, descending
    into content controls (w:sdt) so template body slots are read too."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in container_el.iterchildren():
        if child.tag == W + "p":
            yield "p", Paragraph(child, doc)
        elif child.tag == W + "tbl":
            yield "tbl", Table(child, doc)
        elif child.tag == W + "sdt":
            content = child.find(W + "sdtContent")
            if content is not None:
                for item in _iter_block_items(content, doc):
                    yield item


def _para_text(p) -> str:
    # p.text already joins runs and includes text inside w:hyperlink in python-docx >= 1.0
    return " ".join(p.text.split())


def convert_docx(path: Path) -> Tuple[List[str], List[dict]]:
    _need("docx", "python-docx", ".docx")
    import docx

    doc = docx.Document(str(path))
    lines: List[str] = []
    tables: List[dict] = []
    pidx = 0
    tidx = 0
    for kind, item in _iter_block_items(doc.element.body, doc):
        if kind == "p":
            text = _para_text(item)
            if not text:
                continue
            style = item.style.name if item.style is not None else ""
            if style.lower().startswith("toc"):
                continue
            pidx += 1
            lvl = _heading_level(style)
            marker = " [doc p.%d]" % pidx
            if lvl:
                lines.append("")
                lines.append("#" * lvl + " " + text + marker)
                lines.append("")
            elif style.lower() in ("list paragraph", "list bullet", "list number") or style.lower().startswith("list"):
                lines.append("- " + text + marker)
            elif style.lower() == "caption":
                lines.append("*" + text + "*" + marker)
            else:
                lines.append(text + marker)
                lines.append("")
        else:
            tidx += 1
            grid: List[List[str]] = []
            for row in item.rows:
                cells: List[str] = []
                prev = None
                for c in row.cells:
                    # merged cells repeat the same _tc; keep one copy
                    if prev is not None and c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(" ".join(c.text.split()))
                grid.append(cells)
            grid = [r for r in grid if any(r)]
            if not grid:
                continue
            header, rows = grid[0], grid[1:]
            src = "[doc table %d]" % tidx
            lines.append("")
            lines.extend(md_table(header, rows))
            lines.append(src)
            lines.append("")
            tables.append({"source": src, "header": header, "rows": rows})
    return lines, tables


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------
def _iter_shapes(shapes):
    """Yield every shape, descending into groups (same walk as nw_pptx_helpers)."""
    from pptx.shapes.group import GroupShape

    for sh in shapes:
        yield sh
        if isinstance(sh, GroupShape):
            for sub in _iter_shapes(sh.shapes):
                yield sub


def _slide_title(slide) -> str:
    try:
        from nw_pptx_helpers import slide_title_text  # richer fallbacks when available
        return slide_title_text(slide) or ""
    except Exception:  # noqa: BLE001 - helper absent or slide odd; fall back
        pass
    try:
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            return " ".join(slide.shapes.title.text_frame.text.split())
    except Exception:  # noqa: BLE001
        pass
    return ""


def convert_pptx(path: Path) -> Tuple[List[str], List[dict]]:
    _need("pptx", "python-pptx", ".pptx")
    from pptx import Presentation

    prs = Presentation(str(path))
    lines: List[str] = []
    tables: List[dict] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = _slide_title(slide)
        marker = "[slide %d]" % n
        lines.append("")
        lines.append("## Slide %d: %s %s" % (n, title or "(untitled)", marker))
        lines.append("")
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:  # noqa: BLE001
            title_shape = None
        tcount = 0
        for sh in _iter_shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                tcount += 1
                grid = []
                for row in sh.table.rows:
                    grid.append([" ".join(c.text.split()) for c in row.cells])
                grid = [r for r in grid if any(r)]
                if grid:
                    src = "[slide %d table %d]" % (n, tcount)
                    lines.extend(md_table(grid[0], grid[1:]))
                    lines.append(src)
                    lines.append("")
                    tables.append({"source": src, "header": grid[0], "rows": grid[1:]})
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = " ".join("".join(r.text for r in para.runs).split()) or " ".join(para.text.split())
                    if not text:
                        continue
                    lvl = para.level or 0
                    lines.append("  " * lvl + "- " + text + " " + marker)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                try:
                    ch = sh.chart
                    cats = [str(c) for c in ch.plots[0].categories] if ch.plots else []
                    header = ["Series"] + cats
                    rows = []
                    for s in ch.series:
                        rows.append([s.name] + [_fmt_number(v) for v in s.values])
                    if rows:
                        src = "[slide %d chart]" % n
                        lines.append("")
                        lines.extend(md_table(header, rows))
                        lines.append(src)
                        lines.append("")
                        tables.append({"source": src, "header": header, "rows": rows})
                except Exception:  # noqa: BLE001 - unusual chart shapes are skipped, not fatal
                    lines.append("- (chart) " + marker)
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                lines.append("")
                lines.append("> Notes: %s %s" % (notes, marker))
    return lines, tables


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------
def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def convert_xlsx(path: Path) -> Tuple[List[str], List[dict]]:
    _need("openpyxl", "openpyxl", ".xlsx")
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: List[str] = []
    tables: List[dict] = []
    for ws in wb.worksheets:
        grid: List[List[str]] = []
        max_cols = 0
        total_rows = 0
        for row in ws.iter_rows(values_only=True):
            total_rows += 1
            if len(grid) >= XLSX_ROW_CAP:
                continue
            vals = [_fmt_number(v) for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if not vals:
                continue
            max_cols = max(max_cols, len(vals))
            grid.append(vals)
        if not grid:
            continue
        src = "[%s!A1:%s%d]" % (ws.title, _col_letter(max_cols or 1), len(grid))
        lines.append("")
        lines.append("## Sheet: %s %s" % (ws.title, src))
        lines.append("")
        lines.extend(md_table(grid[0], grid[1:]))
        if total_rows > XLSX_ROW_CAP:
            lines.append("")
            lines.append("_(%d more rows not shown; sheet has %d rows, cap is %d)_" % (total_rows - XLSX_ROW_CAP, total_rows, XLSX_ROW_CAP))
        lines.append("")
        tables.append({"source": src, "header": grid[0], "rows": grid[1:]})
    return lines, tables


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------
def convert_pdf(path: Path) -> Tuple[List[str], List[dict]]:
    _need("pypdf", "pypdf", ".pdf")
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    lines: List[str] = []
    for n, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page should not sink the file
            text = "(text could not be extracted: %s)" % exc
        lines.append("")
        lines.append("## Page %d [page %d]" % (n, n))
        lines.append("")
        for para in re.split(r"\n\s*\n", text.strip()):
            para = " ".join(para.split())
            if para:
                lines.append(para + " [page %d]" % n)
                lines.append("")
    return lines, []


# ---------------------------------------------------------------------------
# vtt
# ---------------------------------------------------------------------------
_TIMING_RE = re.compile(r"^(\d{1,2}:)?\d{2}:\d{2}[.,]\d{3}\s*-->\s*(\d{1,2}:)?\d{2}:\d{2}[.,]\d{3}")
_VOICE_RE = re.compile(r"<v(?:\.[^\s>]*)?\s+([^>]+)>(.*?)(?:</v>|$)", re.DOTALL)
_SPEAKER_RE = re.compile(r"^([A-Z][\w.'\- ]{0,60}?):\s+(.*)$")


def _hms(ts: str) -> str:
    ts = ts.replace(",", ".")
    parts = ts.split(":")
    if len(parts) == 2:
        parts = ["00"] + parts
    h, m, s = parts[0], parts[1], parts[2].split(".")[0]
    return "%02d:%s:%s" % (int(h), m, s)


def convert_vtt(path: Path) -> Tuple[List[str], List[dict]]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    blocks = re.split(r"\r?\n\s*\r?\n", raw.strip())
    cues: List[Tuple[str, str, str]] = []  # (speaker, start, text)
    for block in blocks:
        rows = [r for r in block.splitlines() if r.strip()]
        if not rows or rows[0].startswith(("WEBVTT", "NOTE", "STYLE", "REGION")):
            continue
        ti = next((i for i, r in enumerate(rows) if _TIMING_RE.match(r.strip())), None)
        if ti is None:
            continue
        start = _hms(rows[ti].split("-->")[0].strip().split()[0])
        text = " ".join(r.strip() for r in rows[ti + 1:])
        speaker = ""
        m = _VOICE_RE.search(text)
        if m:
            speaker = m.group(1).strip()
            text = " ".join(_VOICE_RE.sub(lambda mm: mm.group(2), text).split())
        else:
            m2 = _SPEAKER_RE.match(text)
            if m2:
                speaker, text = m2.group(1).strip(), m2.group(2)
        text = re.sub(r"<[^>]+>", "", text).strip()
        if text:
            cues.append((speaker, start, text))
    # Merge consecutive cues from the same speaker into one paragraph, keeping
    # the start time of the first cue in the run.
    merged: List[Tuple[str, str, List[str]]] = []
    for speaker, start, text in cues:
        if merged and merged[-1][0] == speaker:
            merged[-1][2].append(text)
        else:
            merged.append((speaker, start, [text]))
    out: List[str] = []
    for speaker, start, texts in merged:
        out.append("**%s** [%s]: %s" % (speaker or "Unknown speaker", start, " ".join(texts)))
        out.append("")
    return out, []


# ---------------------------------------------------------------------------
# txt / md
# ---------------------------------------------------------------------------
def convert_text(path: Path) -> Tuple[List[str], List[dict]]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    out: List[str] = []
    for i, line in enumerate(raw.splitlines(), start=1):
        if i % TXT_MARK_EVERY == 1:
            out.append(line.rstrip() + " [line %d]" % i)
        else:
            out.append(line.rstrip())
    return out, []


# ---------------------------------------------------------------------------
def header_block(path: Path, ftype: str, label_shown: str) -> List[str]:
    now = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    lines = [
        "# Source: %s" % path.name,
        "- type: %s" % ftype,
        "- size: %d bytes" % path.stat().st_size,
        "- extracted: %s" % now,
        "- label: %s" % label_shown,
    ]
    prov = {
        "docx": "[doc p.N] is a running paragraph index, not a page number (Word files carry no reliable page numbers)",
        "pptx": "[slide N] is the slide number in deck order",
        "xlsx": "[Sheet!A1:F40] is the cell range each table was read from; rows capped at %d per sheet" % XLSX_ROW_CAP,
        "pdf": "[page N] is the PDF page number",
        "vtt": "[HH:MM:SS] is the start time of the first cue merged into each paragraph",
        "txt": "[line N] marks every %dth source line" % TXT_MARK_EVERY,
        "md": "[line N] marks every %dth source line" % TXT_MARK_EVERY,
    }.get(ftype)
    if prov:
        lines.append("- provenance: %s" % prov)
    lines.append("")
    lines.append("---")
    lines.append("")
    return lines


CONVERTERS = {"docx": convert_docx, "pptx": convert_pptx, "xlsx": convert_xlsx,
              "pdf": convert_pdf, "vtt": convert_vtt, "txt": convert_text, "md": convert_text}


def read_source(path: Path, max_chars: Optional[int] = None) -> dict:
    if not path.exists():
        raise ReadError("%s does not exist." % path, "Check the path and try again.")
    ftype = TYPES.get(path.suffix.lower())
    if not ftype:
        raise ReadError("Unsupported file type '%s'." % path.suffix,
                        "Supported: " + ", ".join(sorted(TYPES)))
    label = read_label(path) if ftype in ("docx", "pptx", "xlsx") else None
    shown, label_value = label_line(label)
    body, tables = CONVERTERS[ftype](path)
    body_text = "\n".join(body).strip("\n") + "\n"
    body_text = re.sub(r"\n{3,}", "\n\n", body_text)
    truncated = False
    if max_chars is not None and max_chars > 0 and len(body_text) > max_chars:
        body_text = body_text[:max_chars].rstrip() + "\n\n_[truncated at %d characters; %d in full]_\n" % (max_chars, len(body_text))
        truncated = True
    text = "\n".join(header_block(path, ftype, shown)) + body_text
    return {"type": ftype, "label": label_value, "label_detail": label, "text": text,
            "tables": tables, "truncated": truncated}


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Convert a source file to Markdown with provenance markers.")
    ap.add_argument("path")
    ap.add_argument("--out", help="write the Markdown here instead of stdout")
    ap.add_argument("--tables-json", nargs="?", const="", default=None, metavar="FILE",
                    help="write every table found as JSON (default name: <out or source stem>.tables.json)")
    ap.add_argument("--max-chars", type=int, default=None)
    ap.add_argument("--json", action="store_true", help="print a one-object summary instead of the Markdown")
    args = ap.parse_args(argv)

    path = Path(args.path).expanduser()
    try:
        res = read_source(path, args.max_chars)
    except ReadError as exc:
        if args.json:
            print(json.dumps({"ok": False, "error": str(exc), "hint": exc.hint}))
        else:
            print("error: %s" % exc, file=sys.stderr)
            if exc.hint:
                print(exc.hint, file=sys.stderr)
        return exc.code
    except Exception as exc:  # noqa: BLE001 - corrupt files etc.; keep the contract shape
        if args.json:
            print(json.dumps({"ok": False, "error": "Could not read %s: %s" % (path.name, exc),
                              "hint": "The file may be corrupt, encrypted, or not really a %s." % path.suffix}))
        else:
            print("error: could not read %s: %s" % (path.name, exc), file=sys.stderr)
        return 1

    out_path = None
    if args.out:
        out_path = Path(args.out).expanduser()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(res["text"], encoding="utf-8")
    tables_path = None
    if args.tables_json is not None:
        if args.tables_json:
            tables_path = Path(args.tables_json).expanduser()
        elif out_path is not None:
            tables_path = out_path.with_suffix(".tables.json")
        else:
            tables_path = Path.cwd() / (path.stem + ".tables.json")
        tables_path.parent.mkdir(parents=True, exist_ok=True)
        tables_path.write_text(json.dumps({"tables": res["tables"]}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    if args.json:
        print(json.dumps({
            "ok": True,
            "path": str(path.resolve()),
            "type": res["type"],
            "chars": len(res["text"]),
            "label": res["label"],
            "label_detail": res["label_detail"],
            "tables": len(res["tables"]),
            "truncated": res["truncated"],
            "out": str(out_path.resolve()) if out_path else None,
            "tables_json": str(tables_path.resolve()) if tables_path else None,
        }, indent=2))
    elif out_path is None:
        sys.stdout.write(res["text"])
    else:
        print("wrote %s (%d chars%s)" % (out_path, len(res["text"]), "; tables: %s" % tables_path if tables_path else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
