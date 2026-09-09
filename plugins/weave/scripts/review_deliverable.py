#!/usr/bin/env python3
"""
Score a built Netwoven deliverable against the house rubric (structure,
storyline, evidence, brand, hygiene) on top of the gates that
validate_deliverable.py / nw_pptx_helpers.validate_deck already enforce.

  review_deliverable.py FILE [--quick] [--rubric default|sow|deck] [--apply]
      [--reviewer NAME] [--report out.md] [--client NAME] [--internal] [--json]

Every finding is {check, passed, evidence, hard, severity, fix, area}. Areas
score 0-20 each (high -8, medium -4, low -1 per failing finding, floor 0); a
hard failure caps the total at 59. --quick keeps gates, titles, hygiene,
leakage and metadata. --apply writes FILE.bak, then makes the safe fixes only
(dashes, empty placeholders, closer order, alt text, creator, title, reviewer
property); MSIP_Label_* and docMetadata/LabelInfo.xml are never touched.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import statistics
import sys
import zipfile
from pathlib import Path

from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import nw_docx_helpers as DH  # noqa: E402
import validate_deliverable as V  # noqa: E402

W, norm, para_text, p_style = DH.W, DH.norm, DH.para_text, DH.p_style
DASHES = V.DASHES
DASH_RX = "[%s]" % DASHES
AREAS = ["structure", "storyline", "evidence", "brand", "hygiene"]
DEDUCT = {"high": 8, "medium": 4, "low": 1}
CUSTOM_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
REVIEWER_PROP = "Weave-Reviewed-By"
CHART_URI = 'uri="http://schemas.openxmlformats.org/drawingml/2006/chart"'
# Validator checks this script re-implements with richer evidence and fixes.
SUPERSEDED = {"house style: no em/en dashes in body", "headings <= 12 words",
              "titles are 12 words or fewer", "SOW: no NW Heading / Heading 4 usage"}
MEDIUM_WARNS = ("density", "visuals", "Executive Summary", "fee table", "deliverables table",
                "author/role", "product ticked", "colours", "fonts", "NW Heading 2")
# Conventional headings that may stay topic labels.
TOPIC_OK = {"executive summary", "appendix", "glossary", "next steps", "table of contents", "contents",
            "revision history", "statement of confidentiality", "agenda", "summary"}
VERBS = set("""is are was were be been am has have had do does did will would can could should shall must may
might need needs cut cuts drive drives deliver delivers block blocks keep keeps reach reaches save saves cost
costs take takes make makes show shows give gives lead leads hold holds fall falls rise rises grow grows drop
drops move moves run runs meet meets miss misses beat beats put puts set sets add adds remain remains stay stays
stand stands carry carries cover covers create creates build builds reduce reduces improve improves require
requires expose exposes leave leaves open opens allow allows enable enables explain explains mean means own owns
lack lacks come comes go goes get gets become becomes follow follows start starts begin begins end ends complete
completes use uses work works help helps support supports fix fixes close closes concentrate concentrates outpace
outpaces rebuild rebuilds lock locks reassign reassigns onboard onboards fell rose grew met went came took made
held kept left began became""".split())
FILLER = ["in today's fast-paced digital landscape", "it could perhaps be argued", "may potentially",
          "it is believed that", "click here", "in order to", "it is important to note", "please note that",
          "as previously mentioned", "at the end of the day", "going forward"]
CANDOUR = ["don't tell", "do not tell", "do not share", "don't share", "internal only", "off the record",
           "between us", "the client won't notice", "not for the client"]
QUESTIONS_RX = re.compile(r"^(questions?|discussion|q\s*&\s*a|any questions|thank you)\s*[?!.]*$", re.I)
# percentages, money, standalone numbers of 3+ digits (years 1900-2099 excluded)
NUM_RX = re.compile(r"\d[\d,]*(?:\.\d+)?\s?%|\$\s?\d|\b(?!(?:19|20)\d\d\b)\d{3,}\b|\b\d{1,3},\d{3}\b")
MONTHS = "january|february|march|april|may|june|july|august|september|october|november|december"


# --------------------------------------------------------------------------
# findings
# --------------------------------------------------------------------------
def F(results, check, passed, evidence, area, severity="medium", fix=None, hard=False):
    results.append({"check": check, "passed": bool(passed), "evidence": evidence, "hard": hard,
                    "severity": severity, "fix": None if passed else fix, "area": area})


def area_for(check):
    c = check.lower()
    for area, keys in (("hygiene", ("house style", "placeholder", "highlight", "forbidden", "demo", "embedded ole",
                                    "protection", "words")),
                       ("evidence", ("fee", "amount", "ref field", "bookmark", "chart series", "product",
                                     "deliverables table", "publishdate", "source")),
                       ("storyline", ("content slide has a title", "titles are", "density", "visuals",
                                      "executive summary written", "headings <=")),
                       ("structure", ("heading", "toc", "section", "closer", "appendix", "cover", "slides",
                                      "layout", "confidentiality", "structure", "deck validator"))):
        if any(k in c for k in keys):
            return area
    return "brand"


def wrap_gate(r):
    sev = "high" if r["hard"] else ("medium" if any(k.lower() in r["check"].lower() for k in MEDIUM_WARNS) else "low")
    return {"check": r["check"], "passed": r["passed"], "evidence": r["evidence"], "hard": r["hard"],
            "severity": sev, "fix": None if r["passed"] else "see validate_deliverable.py evidence",
            "area": area_for(r["check"])}


def words(t):
    return re.findall(r"[A-Za-z0-9][A-Za-z0-9'%$.,-]*", t or "")


def has_verb(title):
    """Conservative: any listed verb/auxiliary, an -ed word, or a non-initial -s word counts."""
    toks = [w.lower().strip(".,'") for w in words(title)]
    for i, w in enumerate(toks):
        if w in VERBS or (w.endswith("ed") and len(w) > 4):
            return True
        if i > 0 and len(w) > 3 and w.endswith("s") and not w.endswith(("ss", "us", "is", "ous")):
            return True  # plural noun or 3rd-person verb: lenient on purpose
    return False


def title_case_like(t):
    ws = [w for w in words(t)[1:] if len(w) > 3 and not w.isupper()]
    return len(ws) >= 3 and all(w[0].isupper() for w in ws)


def check_titles(results, titles, rubric, label):
    """Storyline checks shared by the ghost deck and the heading outline (item 2)."""
    over = [t for t in titles if len(words(t)) > 12]
    F(results, "%s: 12 words or fewer" % label, not over, "%d over; e.g. %s" % (len(over), over[:2]) if over else "ok",
      "storyline", "medium", "Shorten to one conclusion of 12 words or fewer")
    colon = [t for t in titles if t.rstrip().endswith(":")]
    F(results, "%s: no trailing colon" % label, not colon, colon[:3] or "ok", "hygiene", "low",
      "Drop the colon; the body follows anyway")
    if rubric == "sow":
        return  # SOW section titles are fixed by the template
    plain = [t for t in titles if not has_verb(t) and t.lower().strip(".:") not in TOPIC_OK]
    topic = [t for t in plain if len(words(t)) <= 3]
    F(results, "%s: no bare topic labels" % label, not topic, topic[:4] or "ok", "storyline", "medium",
      "Rewrite each as the conclusion of the slide/section (subject, verb, object)")
    noverb = [t for t in plain if t not in topic]
    F(results, "%s: contain a finite verb (action titles)" % label, not noverb, noverb[:4] or "ok",
      "storyline", "low", "State what the evidence proves, as a full sentence")
    seen, dups = set(), []
    for t in titles:
        k = t.lower().strip()
        if k in seen:
            dups.append(t)
        seen.add(k)
    F(results, "%s: no duplicates" % label, not dups, dups[:3] or "ok", "structure", "medium",
      "Merge the repeated slides/sections or sharpen one message")
    tc = [t for t in titles if title_case_like(t)]
    F(results, "%s: sentence case" % label, not tc, tc[:3] or "ok", "hygiene", "low",
      "Use sentence case; the style applies caps where the template wants them")


def check_prose(results, lines, kind, date_lines=None):
    """Writing hygiene and leakage over plain text lines (items 8, 9)."""
    dash = [t[:70] for t in lines if re.search(DASH_RX, t)]
    F(results, "hygiene: no em/en dashes", not dash, "%d line(s); e.g. %s" % (len(dash), dash[:2]) if dash else "ok",
      "hygiene", "medium", "Rejoin with a comma or 'to' for ranges (--apply does this)")
    low = "\n".join(lines).lower()
    fill = sorted(set(f for f in FILLER if f in low))
    F(results, "hygiene: no filler phrases", not fill, fill or "ok", "hygiene", "low", "Cut the phrase; keep the fact")
    dbl = [t[:60] for t in lines if re.search(r"\S {2,}\S", t)]
    F(results, "hygiene: no double spaces", not dbl, dbl[:2] or "ok", "hygiene", "low", "Collapse to one space")
    cand = sorted(set(c for c in CANDOUR if c in low))
    F(results, "leakage: no internal candour", not cand, cand or "ok", "hygiene", "high",
      "Remove the passage and move it to the flag list; it must not reach the client")
    # dates: prose only (tables and bound SOW date fields legitimately use M/d/yyyy)
    dl = "\n".join(date_lines if date_lines is not None else lines).lower()
    iso = re.search(r"\b\d{4}-\d{2}-\d{2}\b", dl) is not None
    us = re.search(r"\b\d{1,2}/\d{1,2}/\d{4}\b", dl) is not None and kind != "sow"
    long_ = re.search(r"\b(%s) \d{1,2}, \d{4}" % MONTHS, dl) is not None
    F(results, "consistency: one date format", sum([iso, us, long_]) <= 1,
      "iso=%s us=%s long=%s" % (iso, us, long_), "hygiene", "low", "Write dates as Month D, YYYY in prose")


def check_trackers(results, text, client):
    """Item 7: number drift on the same noun and client-name variants."""
    by_noun = {}
    for num, noun in re.findall(r"\b(\d[\d,]*)\s+([a-z]{3,})\b", text):
        try:
            by_noun.setdefault(noun.rstrip("s"), set()).add(float(num.replace(",", "")))
        except ValueError:
            pass
    drift = []
    for noun, nums in by_noun.items():
        vals = sorted(n for n in nums if n >= 10)
        drift += ["%s: %g vs %g" % (noun, a, b) for a, b in zip(vals, vals[1:]) if a != b and (b - a) / b < 0.1]
    F(results, "consistency: same noun, same number", not drift, drift[:3] or "ok", "evidence", "low",
      "Confirm which figure is right and use it everywhere (only near-equal figures are flagged)")
    if client:
        first = re.escape((words(client) or [client])[0])
        rx = r"\b%s\b(?:[ ,.]{1,2}(?:Inc|Ltd|LLC|Corp|Corporation|Limited|Co|plc|GmbH)\.?)?" % first
        groups = {}
        for f in set(m.strip() for m in re.findall(rx, text, re.I)):
            groups.setdefault(re.sub(r"[^a-z0-9]", "", f.lower()), set()).add(f)
        var = [sorted(g) for g in groups.values() if len(g) > 1]
        F(results, "consistency: client name written one way", not var, var[:2] or "ok", "brand", "low",
          "Use the legal name exactly as the Company field carries it")


# --------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------
def docx_outline(pkg, rubric):
    prefix = "Heading" if rubric == "sow" else "NWHeading"
    out = []
    for p in pkg.body_paragraphs():
        m = re.fullmatch(prefix + r"([1-4])", p_style(p) or "")
        if m and norm(para_text(p)):
            out.append({"level": int(m.group(1)), "title": norm(para_text(p)), "style": p_style(p)})
    return out


def sow_skip(pkg, body_ps, tables=False):
    """Paragraphs whose dashes are template boilerplate (validator rule; tables too for --apply)."""
    skip, in_gloss = set(), False
    for p in body_ps:
        if p_style(p) == "Heading1":
            in_gloss = norm(para_text(p)).lower() == "glossary"
        if in_gloss or norm(V.own_text(p)).startswith("Total Project Estimate"):
            skip.add(id(p))
    if tables:
        skip |= {id(p) for tbl in pkg.body_tables() for p in tbl.iter(W + "p")}
    return skip


def review_docx(path, kind, rubric, args, results, ghost):
    for r in V.check_docx(path, kind, args.internal, args.client):
        if r["check"] not in SUPERSEDED:
            results.append(wrap_gate(r))
    pkg = V.Pkg(path)
    body_ps = pkg.body_paragraphs()
    outline = docx_outline(pkg, rubric)
    ghost.extend(outline)
    check_titles(results, [o["title"] for o in outline if o["level"] <= 2], rubric, "headings")
    skip = sow_skip(pkg, body_ps) if kind == "sow" else set()
    lines = [norm(V.own_text(p)) for p in body_ps if id(p) not in skip and norm(V.own_text(p))]
    top_lines = [norm(V.own_text(el)) for el in pkg.body_children() if el.tag == W + "p"]
    check_prose(results, lines, kind, top_lines)
    if rubric == "sow":
        bad = [o["title"] for o in outline if o["level"] == 4]
        bad += sorted(set(p_style(p) for p in body_ps if (p_style(p) or "").startswith("NWHeading")))
        F(results, "structure: SOW uses Heading 1-3 only (no Heading 4 / NW Heading)", not bad, bad[:3] or "none",
          "structure", "high", "Fold level-4 text into its level-3 parent; restyle NW Heading n as Heading n", hard=True)
    if args.quick:
        return
    # item 5: MECE (lonely children), section length outliers
    lonely = []
    for i, o in enumerate(outline):
        kids = 0
        for n in outline[i + 1:]:
            if n["level"] <= o["level"]:
                break
            kids += n["level"] == o["level"] + 1
        if kids == 1:
            lonely.append(o["title"])
    F(results, "structure: no heading with a single child (MECE)", not lonely, lonely[:3] or "ok", "structure",
      "medium", "A group of one is not a group: promote the child or add its siblings")
    top, counts, cur = (outline[0]["style"] if outline else None), [], 0
    for p in body_ps:
        if p_style(p) == top:
            counts.append(cur)
            cur = 0
        elif norm(para_text(p)):
            cur += 1
    counts = [c for c in counts[1:] + [cur] if c] if counts else []
    med = statistics.median(counts) if counts else 0
    # 4x median and 15+ paragraphs (kept loose); template-fixed SOW sections are exempt
    outl = [c for c in counts if med and c > 4 * med and c > 15 and rubric != "sow"]
    F(results, "structure: section lengths balanced", not outl, "paragraphs per section %s" % counts,
      "structure", "low", "Split the long section or move detail to an appendix")
    # item 6: numbers in tables / captioned figures need a source line (not for SOW: the SOW is the source)
    if rubric != "sow":
        unsourced, children = [], list(pkg.body_children())
        for i, el in enumerate(children):
            near = [c for c in children[max(0, i - 2):i + 3] if c is not el and c.tag == W + "p"]
            is_tbl = el.tag == W + "tbl"
            is_fig = el.tag == W + "p" and el.find(".//" + DH.PIC + "pic") is not None and \
                any(p_style(c) == "Caption" for c in near)
            txt = norm(para_text(el))
            if not (is_tbl or is_fig) or (is_tbl and txt.startswith("Date")):  # revision history
                continue
            if (is_fig or NUM_RX.search(txt)) and not any(norm(para_text(c)).lower().startswith(("source", "note")) for c in near):
                unsourced.append(("table " if is_tbl else "figure ") + txt[:40])
        F(results, "evidence: numbers carry a source line", not unsourced, unsourced[:3] or "ok", "evidence",
          "medium", "Add a 'Source: ...' paragraph under the table or figure")
    check_trackers(results, "\n".join(lines), args.client)


# --------------------------------------------------------------------------
# pptx
# --------------------------------------------------------------------------
def deck_ghost(prs, H):
    """[{n, title, layout, evidence, titled, role}] for every slide."""
    out = []
    for i, s in enumerate(prs.slides, start=1):
        lname = H.layout_name(s)
        role = ("cover" if i == 1 or lname == "Title Slide for Verticals" else "closer" if H.closer_kind(s)
                else "section" if lname == "Section Header" else "body")
        xml = etree.tostring(s._element).decode("utf-8", "replace")
        ev = [k for k, hit in (("chart", CHART_URI in xml), ("table", "<a:tbl>" in xml),
                               ("picture", "<p:pic" in xml), ("stats", lname == "Three Stat")) if hit]
        try:
            titled = H.layout_has_title(s.slide_layout)
        except Exception:  # noqa: BLE001 - detached layout
            titled = False
        out.append({"n": i, "title": H.slide_title_text(s), "layout": lname, "evidence": ev, "titled": titled, "role": role})
    return out


def review_pptx(path, args, results, ghost):
    import pptx
    from pptx.util import Inches
    import nw_pptx_helpers as H
    prs = pptx.Presentation(str(path))
    for r in H.validate_deck(prs, str(path), args):
        if r["check"] not in SUPERSEDED:
            results.append(wrap_gate(r))
    entries, slides = deck_ghost(prs, H), list(prs.slides)
    body = [e for e in entries if e["role"] == "body"]
    ghost.extend(e for e in entries if e["role"] in ("body", "section"))
    check_titles(results, [e["title"] for e in body if e["titled"] and e["title"]], "deck", "slide titles")
    q = ["slide %d %r" % (e["n"], e["title"]) for e in body if QUESTIONS_RX.match(e["title"] or "")]
    F(results, "structure: no Questions/Discussion slide", not q, q or "ok", "structure", "medium",
      "End with a next-steps slide; the template closers already thank the audience")
    lines = [t for e in entries if e["role"] != "closer" for t in H.slide_text(slides[e["n"] - 1]).split("\n") if t.strip()]
    check_prose(results, lines, "deck")
    if args.quick:
        return

    def body_paras(s):
        tshape = H._title_shape(s)
        return [p.text for sh in H._iter_shapes(s.shapes) if getattr(sh, "has_text_frame", False) and sh.has_text_frame
                and (tshape is None or sh._element is not tshape._element) for p in sh.text_frame.paragraphs
                if p.text.strip() and not p.text.strip().lower().startswith("source")]

    # item 3: horizontal logic (executive summary statements vs sections / evidence slides)
    exec_e = next((e for e in body if "summary" in (e["title"] or "").lower()), None)
    if exec_e is None and body and body[0]["layout"] in ("Title and Content", "Three Column"):
        exec_e = body[0]
    if exec_e is not None:
        n_stat = 3 if exec_e["layout"] == "Three Column" else len(body_paras(slides[exec_e["n"] - 1]))
        sections = [e for e in entries if e["role"] == "section"]
        target = len(sections) if len(sections) >= 2 else len([e for e in body if e["evidence"] and e is not exec_e])
        F(results, "storyline: executive summary mirrors the sections", n_stat == target,
          "slide %d has %d statements vs %d %s" % (exec_e["n"], n_stat, target, "sections" if len(sections) >= 2 else "evidence slides"),
          "storyline", "low", "One key-line statement per section, in the same order (deck-storyline.md section 6)")
    else:
        F(results, "storyline: executive summary slide present", False, "no summary slide after the cover",
          "storyline", "medium", "Add a Title and Content slide whose title is the governing thought")
    # item 4 + 6: one message per slide (deck-storyline.md section 8) and sourced numbers
    dense, twocharts, unsourced = [], [], []
    for e in body:
        s = slides[e["n"] - 1]
        paras = body_paras(s)
        nw = sum(len(words(p)) for p in paras)
        if nw > 75 or (len(paras) > 6 and e["layout"] != "Three Stat"):
            dense.append("slide %d (%d lines, %d words)" % (e["n"], len(paras), nw))
        if etree.tostring(s._element).decode("utf-8", "replace").count(CHART_URI) >= 2:
            twocharts.append("slide %d" % e["n"])
        needs = any(k in e["evidence"] for k in ("chart", "table", "stats")) or NUM_RX.search("\n".join(paras))
        if e is not exec_e and needs:
            has_src = any(t.strip().lower().startswith(("source", "note")) for t in H.slide_text(s).split("\n"))
            near = any(sh.top is not None and sh.top >= H.SOURCE_LINE_TOP - Inches(0.3) and H.shape_text(sh).strip()
                       for sh in s.shapes)
            if not (has_src or near):
                unsourced.append("slide %d" % e["n"])
    F(results, "storyline: one message per slide (<= 75 words, <= 6 lines)", not dense, dense[:4] or "ok",
      "storyline", "low", "Split the slide or cut bullets that restate the title")
    F(results, "storyline: one chart per slide", not twocharts, twocharts or "ok", "storyline", "medium",
      "One message, one proof: move the second chart to its own slide")
    F(results, "evidence: numbers carry a source line", not unsourced, unsourced[:4] or "ok", "evidence", "medium",
      "add_source_line(slide, 'Source: ...') under the visual")
    check_trackers(results, "\n".join(lines), args.client)


# --------------------------------------------------------------------------
# metadata (item 10) and --apply
# --------------------------------------------------------------------------
def core_tag(text, tag):
    """Read one docProps/core.xml element from raw XML text: '' for an empty
    or self-closing tag (attributes and all, matching DH.scrub_metadata's own
    pattern), None if the tag is not present at all."""
    m = re.search(r"<%s\b[^>]*>(.*?)</%s>" % (tag, tag), text, re.S)
    if m:
        return norm(m.group(1))
    return "" if re.search(r"<%s\b[^>]*/>" % tag, text) else None


def check_metadata(path, results):
    with zipfile.ZipFile(str(path)) as z:
        part = lambda n: z.read(n).decode("utf-8", "replace") if n in z.namelist() else ""  # noqa: E731
        core, custom, doc = part("docProps/core.xml"), part("docProps/custom.xml"), part("word/document.xml")
    creator, title = core_tag(core, "dc:creator") or "", core_tag(core, "dc:title") or ""
    F(results, "metadata: creator scrubbed", creator in ("", "Netwoven"), "creator=%r" % creator, "brand", "medium",
      "Set dc:creator to Netwoven (--apply does this)")
    F(results, "metadata: dc:title set", bool(title), "title=%r" % title, "brand", "medium",
      "Set the document title (--apply copies the first title)")
    weave = sorted(set(re.findall(r'name="(Weave_[A-Za-z0-9]+)"', custom)))
    F(results, "metadata: Weave_* properties (informational)", True, weave or "absent (built without a spec)", "brand", "low")
    rev_by = 'name="%s"' % REVIEWER_PROP in custom
    rev_row = re.search(r"reviewed by\s+[A-Z][a-z]+", re.sub(r"<[^>]+>", "", doc)) is not None
    F(results, "metadata: reviewer recorded", rev_by or rev_row, "%s property %s; revision row %s" %
      (REVIEWER_PROP, "present" if rev_by else "absent", "found" if rev_row else "not found"), "brand", "low",
      "Run --apply --reviewer NAME or add a 'Reviewed by NAME' revision row")


def rewrite_zip(path, replace):
    """Rewrite members in place (adds missing ones); untouched parts such as LabelInfo.xml are copied byte for byte."""
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(str(path)) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        done = set()
        for item in zin.infolist():
            done.add(item.filename)
            zout.writestr(item, replace.get(item.filename, zin.read(item.filename)))
        for name, data in replace.items():
            if name not in done:
                zout.writestr(name, data)
    os.replace(tmp, str(path))


def fix_dashes(nodes):
    """House-style dash replacement over the text nodes of one paragraph. Returns changed node count."""
    n = 0
    for i, node in enumerate(nodes):
        t = node.text or ""
        if not re.search(DASH_RX, t):
            continue
        nxt = (nodes[i + 1].text or "") if i + 1 < len(nodes) else ""
        t = re.sub(r"(\d)\s*%s\s*(\d)" % DASH_RX, r"\1 to \2", t)  # numeric range
        if i == 0:
            t = re.sub(r"^\s*%s\s*" % DASH_RX, "", t)  # manual bullet
        t = re.sub(r"\s*%s\s*$" % DASH_RX, "," if nxt.startswith(" ") else ", ", t)  # run boundary
        node.text = re.sub(r"\s*%s\s*" % DASH_RX, ", ", t)
        if node.text != node.text.strip():
            node.set(DH.XML_SPACE, "preserve")
        n += 1
    return n


def set_custom_property(path, name, value):
    """Add/update one docProps/custom.xml property; creates the part (plus content type and rel) when absent."""
    with zipfile.ZipFile(str(path)) as z:
        raw = z.read("docProps/custom.xml") if "docProps/custom.xml" in z.namelist() else None
        ct, rels = z.read("[Content_Types].xml").decode("utf-8"), z.read("_rels/.rels").decode("utf-8")
    replace = {}
    if raw is None:
        root = etree.Element("{%s}Properties" % CUSTOM_NS, nsmap={None: CUSTOM_NS, "vt": VT_NS})
        if "/docProps/custom.xml" not in ct:
            replace["[Content_Types].xml"] = ct.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType='
                                                        '"application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>').encode("utf-8")
        if "relationships/custom-properties" not in rels:
            rid = max([int(x) for x in re.findall(r'Id="rId(\d+)"', rels)] or [0]) + 1
            replace["_rels/.rels"] = rels.replace("</Relationships>", '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/'
                                                  'officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>' % rid).encode("utf-8")
    else:
        root = etree.fromstring(raw)
    props = root.findall("{%s}property" % CUSTOM_NS)
    prop = next((p for p in props if p.get("name") == name), None)
    if prop is None:
        pid = max([int(p.get("pid", "1")) for p in props] + [1]) + 1
        prop = etree.SubElement(root, "{%s}property" % CUSTOM_NS, fmtid=FMTID, pid=str(pid), name=name)
        etree.SubElement(prop, "{%s}lpwstr" % VT_NS)
    prop[0].text = value
    replace["docProps/custom.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    rewrite_zip(path, replace)


def apply_docx(path, kind, applied):
    pkg = V.Pkg(path)
    body_ps = pkg.body_paragraphs()
    skip = sow_skip(pkg, body_ps, tables=True) if kind == "sow" else set()
    n = sum(fix_dashes(list(p.iter(W + "t"))) for p in body_ps if id(p) not in skip)
    alt = 0
    for docPr in pkg.doc.iter(DH.WP + "docPr"):
        drawing = docPr.getparent().getparent()
        if drawing is not None and drawing.find(".//" + DH.PIC + "pic") is not None and not docPr.get("descr"):
            docPr.set("descr", "Figure")
            alt += 1
    replace = {}
    if n or alt:
        replace["word/document.xml"] = etree.tostring(pkg.doc, xml_declaration=True, encoding="UTF-8", standalone=True)
        applied += (["replaced dashes in %d run(s)" % n] if n else []) + (["alt text 'Figure' on %d picture(s)" % alt] if alt else [])
    core = new = pkg.read("docProps/core.xml")
    m = re.search(r"<dc:title>(.*?)</dc:title>|<dc:title/>", new, re.S)
    first = next((norm(para_text(p)) for p in body_ps if re.fullmatch(r"(NW)?Heading1", p_style(p) or "")), "")
    if m and not norm(m.group(1) or "") and first:
        new = new.replace(m.group(0), "<dc:title>%s</dc:title>" % first.replace("&", "&amp;").replace("<", "&lt;"))
        applied.append("dc:title set to %r" % first)
    if new != core:
        replace["docProps/core.xml"] = new.encode("utf-8")
    pkg.z.close()
    if replace:
        rewrite_zip(path, replace)
    # Canonical metadata scrub (creator/lastModifiedBy -> Netwoven, including
    # self-closing tags; cp:lastPrinted dropped; dcterms:created/modified ->
    # now; cp:revision -> 1). docProps/custom.xml (MSIP_Label_*) and
    # docMetadata/LabelInfo.xml are outside DH.scrub_metadata's reach.
    before = new
    DH.scrub_metadata(str(path))
    with zipfile.ZipFile(str(path)) as z:
        after = z.read("docProps/core.xml").decode("utf-8", "replace")
    for t in ("dc:creator", "cp:lastModifiedBy"):
        if core_tag(before, t) != core_tag(after, t):
            applied.append("%s scrubbed to Netwoven" % t)
    if "<cp:lastPrinted" in before and "<cp:lastPrinted" not in after:
        applied.append("cp:lastPrinted cleared")
    if core_tag(before, "cp:revision") not in (None, "1") and core_tag(after, "cp:revision") == "1":
        applied.append("cp:revision reset to 1")


def apply_pptx(path, applied):
    import pptx
    from pptx.oxml.ns import qn
    import nw_pptx_helpers as H
    prs = pptx.Presentation(str(path))
    n = alt = 0
    for s in prs.slides:
        if H.closer_kind(s):
            continue
        for sh in H._iter_shapes(s.shapes):
            frames = [sh.text_frame] if getattr(sh, "has_text_frame", False) and sh.has_text_frame else []
            if getattr(sh, "has_table", False) and sh.has_table:
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
            n += sum(fix_dashes(list(p._p.iter(qn("a:t")))) for tf in frames for p in tf.paragraphs)
            if sh.shape_type == 13:  # PICTURE
                c = sh._element.find(".//" + qn("p:cNvPr"))
                if c is not None and not c.get("descr"):
                    H.set_alt_text(sh, "Figure")
                    alt += 1
    applied += (["replaced dashes in %d run(s)" % n] if n else []) + (["alt text 'Figure' on %d picture(s)" % alt] if alt else [])
    cp = prs.core_properties
    for attr in ("author", "last_modified_by"):
        if (getattr(cp, attr) or "").strip() not in ("", "Netwoven"):
            setattr(cp, attr, "Netwoven")
            applied.append("%s scrubbed to Netwoven" % attr)
    if not (cp.title or "").strip() and len(prs.slides):
        cp.title = H.slide_title_text(list(prs.slides)[0])
        applied.append("dc:title set to %r" % cp.title)
    kinds = [H.closer_kind(s) for s in prs.slides]
    tail = [k for k in kinds if k]
    empties = sum(1 for s in prs.slides for sh in s.shapes if H.is_empty_placeholder(sh))
    # tmp + move, not prs.save(path) directly -- see new_deliverable_docx.py's
    # finish_document() for why; --apply is a rewrite of the caller's own
    # file, so this is exactly the path a crash-mid-save must not corrupt.
    tmp = str(path) + ".tmp"
    prs.save(tmp)
    shutil.move(tmp, str(path))
    try:
        from new_deck_pptx import finalize_deck
        finalize_deck(str(path), prune_empty=True, log=lambda *a, **k: None)
    except ImportError:  # minimal re-implementation of the finalize step
        prs = pptx.Presentation(str(path))
        H.move_closers_last(prs)
        for s in prs.slides:
            H.remove_empty_placeholders(s)
        tmp2 = str(path) + ".tmp"
        prs.save(tmp2)
        shutil.move(tmp2, str(path))
        H.patch_app_xml(str(path))
    if empties:
        applied.append("removed %d empty placeholder(s)" % empties)
    if tail and kinds[-len(tail):] != tail:
        applied.append("moved closers to the tail")


# --------------------------------------------------------------------------
def score(findings):
    areas = {a: 20 for a in AREAS}
    for f in findings:
        if not f["passed"]:
            areas[f["area"]] = max(0, areas[f["area"]] - DEDUCT[f["severity"]])
    total = sum(areas.values())
    if any(f["hard"] and not f["passed"] for f in findings):
        total = min(total, 59)
    verdict = ("ready to send after the checklist" if total >= 85 else
               "fix the listed items" if total >= 70 else "not ready")
    return total, verdict, areas


def review(path, args):
    path = Path(path)
    kind = V.detect_type(path)
    rubric = args.rubric if args.rubric and args.rubric != "default" else {"sow": "sow", "deck": "deck"}.get(kind, "default")
    applied = []
    if args.apply:
        shutil.copy2(str(path), str(path) + ".bak")
        applied.append("backup written to %s.bak" % path.name)
        if kind == "deck":
            apply_pptx(path, applied)
        else:
            apply_docx(path, kind, applied)
        if args.reviewer:
            set_custom_property(path, REVIEWER_PROP, args.reviewer)
            applied.append("%s = %r" % (REVIEWER_PROP, args.reviewer))
    results, ghost = [], []
    if kind == "deck":
        review_pptx(path, args, results, ghost)
    else:
        review_docx(path, kind, rubric, args, results, ghost)
    check_metadata(path, results)
    total, verdict, areas = score(results)
    return {"ok": True, "file": str(path), "kind": kind, "rubric": rubric, "score": total, "verdict": verdict,
            "areas": areas, "findings": results, "ghost": ghost, "applied": applied}


def render_report(rep):
    out = ["# Review: %s" % Path(rep["file"]).name, "",
           "**Verdict:** %s (score %d/100, %s rubric)" % (rep["verdict"], rep["score"], rep["rubric"]), "",
           "| Area | Score |", "|---|---|"] + ["| %s | %d/20 |" % (a, rep["areas"][a]) for a in AREAS]
    out += ["", "## %s" % ("Ghost deck" if rep["kind"] == "deck" else "Heading outline"), ""]
    for g in rep["ghost"]:
        if rep["kind"] == "deck":
            out.append("%02d  %s%s" % (g["n"], g["title"] or "(untitled)", "  [%s]" % ", ".join(g["evidence"]) if g["evidence"] else ""))
            out.append("    layout: %s%s" % (g["layout"], "  (section divider)" if g["role"] == "section" else ""))
        else:
            out.append("%s- %s" % ("  " * (g["level"] - 1), g["title"]))
    fails = [f for f in rep["findings"] if not f["passed"]]
    out += ["", "## Findings (%d)" % len(fails), ""]
    for sev in ("high", "medium", "low"):
        group = [f for f in fails if f["severity"] == sev]
        if group:
            out.append("### %s" % sev.capitalize())
            for f in group:
                out.append("- **%s** [%s%s]: %s" % (f["check"], f["area"], ", hard" if f["hard"] else "", f["evidence"]))
                if f["fix"]:
                    out.append("  - fix: %s" % f["fix"])
            out.append("")
    if rep["applied"]:
        out += ["## Applied", ""] + ["- %s" % a for a in rep["applied"]] + [""]
    out += ["## Checked but fine", ""] + ["- %s" % f["check"] for f in rep["findings"] if f["passed"]]
    return "\n".join(out) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--quick", action="store_true")
    ap.add_argument("--rubric", choices=["default", "sow", "deck"], default=None)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--reviewer", default=None, help="record %s (with --apply)" % REVIEWER_PROP)
    ap.add_argument("--report", default=None)
    ap.add_argument("--client", default=None)
    ap.add_argument("--internal", action="store_true")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args(argv)
    if not Path(args.file).exists():
        print(json.dumps({"ok": False, "error": "%s not found" % args.file, "hint": "pass a built .docx or .pptx"}))
        return 1
    try:
        rep = review(args.file, args)
    except Exception as exc:  # noqa: BLE001 - one JSON error object on any failure
        print(json.dumps({"ok": False, "error": "%s: %s" % (type(exc).__name__, exc), "hint": "run validate first"}))
        return 1
    if args.report:
        Path(args.report).write_text(render_report(rep), encoding="utf-8")
    print(json.dumps(rep, indent=2) if args.json else render_report(rep))
    return 0


if __name__ == "__main__":
    sys.exit(main())
